"""Build the reading guide from its reviewed Markdown source."""
from pathlib import Path
import re
import html
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).absolute().parents[1]
OUT = ROOT / 'deliverables/soc_top_view'
QA = ROOT / 'build/wave_guide_document'
QA.mkdir(parents=True, exist_ok=True)
source = (OUT / '波形阅读指南.md').read_text(encoding='utf-8')
doc = Document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21)
sec.top_margin = sec.bottom_margin = Cm(1.6)
sec.left_margin = sec.right_margin = Cm(1.8)
for name, size in [('Normal', 10), ('Title', 21), ('Heading 1', 16), ('Heading 2', 12)]:
    s = doc.styles[name]
    s.font.name = 'Microsoft YaHei'
    s.font.size = Pt(size)
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    s.paragraph_format.space_after = Pt(6)
    s.paragraph_format.line_spacing = 1.12

body = []
lines = source.splitlines()
i = 0
while i < len(lines):
    line = lines[i]
    i += 1
    if not line.strip():
        continue
    if line == '---':
        doc.add_page_break()
        body.append('<div class="pagebreak"></div>')
    elif line.startswith('#'):
        level = len(line) - len(line.lstrip('#'))
        text = line[level:].strip()
        doc.add_paragraph(text, ['Title', 'Heading 1', 'Heading 2'][level-1])
        body.append(f'<h{level}>{html.escape(text)}</h{level}>')
    elif line.startswith('|'):
        rows = [line]
        while i < len(lines) and lines[i].startswith('|'):
            rows.append(lines[i]); i += 1
        data = [[c.strip() for c in r.strip('|').split('|')] for r in rows if not re.match(r'^\|[- :|]+\|$', r)]
        table = doc.add_table(rows=0, cols=len(data[0]))
        table.autofit = False
        widths = [7.2, 10.2] if len(data[0]) == 2 else [5.7, 3.3, 8.4]
        if len(data[0]) == 3 and '事件' in data[0][0]: widths = [3.4, 4.3, 9.7]
        for col, width in zip(table.columns, widths): col.width = Cm(width)
        borders = OxmlElement('w:tblBorders')
        for edge in ['top','left','bottom','right','insideH','insideV']:
            el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
        table._tbl.tblPr.append(borders)
        hrows=[]
        for r, row in enumerate(data):
            cells=table.add_row().cells
            for c, text in enumerate(row):
                cells[c].text=text
                tcpr=cells[c]._tc.get_or_add_tcPr()
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'), 'E7EEF5' if r==0 else 'FFFFFF');tcpr.append(shade)
                margins=OxmlElement('w:tcMar')
                for edge in ['top','bottom','left','right']:
                    el=OxmlElement('w:'+edge);el.set(qn('w:w'),'70');el.set(qn('w:type'),'dxa');margins.append(el)
                tcpr.append(margins)
                for p in cells[c].paragraphs:
                    p.paragraph_format.space_after=Pt(2)
                    for run in p.runs: run.font.size=Pt(9);run.bold=(r==0)
            if r==0:
                el=OxmlElement('w:tblHeader');table.rows[r]._tr.get_or_add_trPr().append(el)
            tag='th' if r==0 else 'td'
            hrows.append('<tr>'+''.join(f'<{tag}>{html.escape(x)}</{tag}>' for x in row)+'</tr>')
        body.append('<table>'+''.join(hrows)+'</table>')
    else:
        doc.add_paragraph(line)
        body.append('<p>'+html.escape(line)+'</p>')
doc.save(QA / '波形阅读指南.docx')
css='''body{font:16px/1.85 "Microsoft YaHei",sans-serif;color:#243244;background:#f3f5f7;margin:0}main{max-width:1040px;margin:32px auto;padding:40px 52px;background:white;border-radius:10px}h1{font-size:30px;color:#000}h2{margin-top:30px;color:#000;font-size:25px}h3{color:#000;font-size:19px}table{border-collapse:collapse;width:100%;margin:18px 0;font-size:15px}td,th{border:1px solid #d9d9d9;padding:10px 14px;text-align:left;overflow-wrap:anywhere}th{background:#e7eef5}a{color:#1655a5}.pagebreak{margin:48px 0}nav{background:#edf3f9;padding:14px 20px}@media print{body{background:#fff;font-size:10pt}main{margin:0;padding:0;max-width:none}.pagebreak{break-before:page;margin:0}h2,h3{break-after:avoid}tr{break-inside:avoid}nav{display:none}}'''
nav='<nav>从状态到接口逐步阅读 · <a href="信号与状态说明.md">61 根信号速查</a> · <a href="soc_architecture.drawio">硬件连接图</a></nav>'
(OUT/'波形阅读指南.html').write_text('<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>六模块 SoC 波形阅读指南</title><style>'+css+'</style><main>'+nav+''.join(body)+'</main></html>',encoding='utf-8')
print('Built reading guide HTML and internal DOCX draft')
